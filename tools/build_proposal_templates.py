from pathlib import Path
from copy import deepcopy
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/workspace/scratch/1cddbdd104f0')
OUT = ROOT / 'outputs' / 'propostas_padrao'
OUT.mkdir(parents=True, exist_ok=True)

PLUGIN_WEBP = Path('/tmp/plugin_logo.webp')
PLUGIN_PNG = OUT / 'plugin_logo.png'
LED_LOGO = Path('/tmp/led_logos/ledcolorido (2).png')

Image.open(PLUGIN_WEBP).convert('RGBA').save(PLUGIN_PNG)

INK = '171717'
MUTED = '666666'
LIGHT = 'F4F4F4'
WHITE = 'FFFFFF'
BLACK = '000000'
ORANGE = 'F25D07'
LED_CYAN = '00D8E8'
LED_MAGENTA = 'EE176A'
LED_BLUE = '2D4EFF'
BORDER = 'D9D9D9'

CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {'top': 90, 'bottom': 90, 'start': 120, 'end': 120}


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for name, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{name}'))
        if node is None:
            node = OxmlElement(f'w:{name}')
            tcMar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def shade(element, fill):
    pPr = element.get_or_add_pPr() if hasattr(element, 'get_or_add_pPr') else element
    shd = pPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        pPr.append(shd)
    shd.set(qn('w:fill'), fill)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def border_cell(cell, color=BORDER, size='6'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = borders.find(qn(f'w:{edge}'))
        if tag is None:
            tag = OxmlElement(f'w:{edge}')
            borders.append(tag)
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), size)
        tag.set(qn('w:color'), color)


def set_run(run, size=10.5, color=INK, bold=False, italic=False, font='Poppins'):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), font)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    return run


def set_paragraph(p, before=0, after=8, line=1.333, alignment=WD_ALIGN_PARAGRAPH.LEFT, keep=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.alignment = alignment
    if keep:
        pf.keep_with_next = True
    return p


def add_text(doc, text, size=10.5, color=INK, bold=False, italic=False,
             before=0, after=8, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.333,
             keep=False):
    p = doc.add_paragraph()
    set_paragraph(p, before, after, line, align, keep)
    set_run(p.add_run(text), size, color, bold, italic)
    return p


def add_heading(doc, text, level, accent):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.text = ''
    sizes = {1: 16, 2: 13, 3: 12}
    before = {1: 18, 2: 12, 3: 8}
    after = {1: 10, 2: 6, 3: 4}
    set_paragraph(p, before[level], after[level], 1.0, keep=True)
    set_run(p.add_run(text), sizes[level], accent if level < 3 else INK, True)
    return p


def add_page_break(doc):
    doc.add_page_break()


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(total))
    tblW.set(qn('w:type'), 'dxa')
    tblInd = tblPr.find(qn('w:tblInd'))
    if tblInd is None:
        tblInd = OxmlElement('w:tblInd')
        tblPr.append(tblInd)
    tblInd.set(qn('w:w'), str(TABLE_INDENT_DXA))
    tblInd.set(qn('w:type'), 'dxa')
    layout = tblPr.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        tblPr.append(layout)
    layout.set(qn('w:type'), 'fixed')
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa)-1)]
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:w'), str(width))
            tcW.set(qn('w:type'), 'dxa')
            set_cell_margins(cell, **CELL_MARGINS)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_cell_text(cell, size=9.5, color=INK, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    for p in cell.paragraphs:
        set_paragraph(p, 0, 0, 1.15, align)
        for run in p.runs:
            set_run(run, size, color, bold)


def add_data_table(doc, headers, rows, widths, accent, total_row=None):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade_cell(cell, accent)
        border_cell(cell, accent)
        style_cell_text(cell, 9.3, WHITE, True, WD_ALIGN_PARAGRAPH.CENTER)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            shade_cell(cells[i], WHITE if r_idx % 2 == 0 else LIGHT)
            border_cell(cells[i])
            align = WD_ALIGN_PARAGRAPH.CENTER if i >= len(row)-2 and len(row) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            style_cell_text(cells[i], 9.3, INK, False, align)
    if total_row:
        cells = table.add_row().cells
        for i, value in enumerate(total_row):
            cells[i].text = value
            shade_cell(cells[i], 'EDEDED')
            border_cell(cells[i], accent, '8')
            style_cell_text(cells[i], 10, INK, True,
                            WD_ALIGN_PARAGRAPH.RIGHT if i == len(total_row)-1 else WD_ALIGN_PARAGRAPH.LEFT)
    p = doc.add_paragraph()
    set_paragraph(p, 0, 5, 1.0)
    return table


def add_label_value_table(doc, rows, accent, widths=(2700, 6660)):
    table = doc.add_table(rows=0, cols=2)
    for idx, (label, value) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        shade_cell(cells[0], LIGHT)
        shade_cell(cells[1], WHITE)
        for cell in cells:
            border_cell(cell)
        style_cell_text(cells[0], 9.3, accent, True)
        style_cell_text(cells[1], 9.5, INK, False)
    set_table_geometry(table, list(widths))
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_callout(doc, label, text, accent):
    p = doc.add_paragraph()
    set_paragraph(p, 5, 10, 1.25)
    shade(p._p, LIGHT)
    set_run(p.add_run(label + ' '), 10.2, accent, True)
    set_run(p.add_run(text), 10.2, INK)
    return p


def add_bullets(doc, items, accent):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        set_paragraph(p, 0, 4, 1.208)
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.194)
        p.text = ''
        set_run(p.add_run(item), 10.2, INK)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.extend([fldChar1, instrText, fldChar2])
    set_run(run, 8.5, MUTED)


def configure_doc(doc, brand_name, accent):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Poppins'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Poppins')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Poppins')
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for level, size, before, after in [(1,16,18,10),(2,13,12,6),(3,12,8,4)]:
        style = styles[f'Heading {level}']
        style.font.name = 'Poppins'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Poppins')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Poppins')
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(accent if level < 3 else INK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.text = ''
    set_paragraph(hp, 0, 0, 1.0)
    set_run(hp.add_run(brand_name.upper() + '  |  PROPOSTA COMERCIAL'), 8.2, MUTED, True)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ''
    set_run(fp.add_run('{{empresa_site}}  |  {{empresa_telefone}}  |  Página '), 8.2, MUTED)
    add_page_number(fp)

    props = doc.core_properties
    props.title = f'Modelo de Proposta Comercial - {brand_name}'
    props.subject = 'Modelo padronizado para geração automática no aplicativo'
    props.author = 'Patrícia'
    props.keywords = 'proposta comercial, template, aplicativo'


def add_cover(doc, logo_path, brand_name, subtitle, accent, logo_width):
    p = doc.add_paragraph()
    set_paragraph(p, 0, 0, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    shade(p._p, BLACK)
    run = p.add_run()
    run.add_picture(str(logo_path), width=Inches(logo_width))
    for _ in range(2):
        spacer = doc.add_paragraph()
        set_paragraph(spacer, 0, 0, 1.0)
        shade(spacer._p, BLACK)
        spacer.add_run(' ')

    title = doc.add_paragraph()
    set_paragraph(title, 0, 4, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
    shade(title._p, BLACK)
    set_run(title.add_run('PROPOSTA COMERCIAL'), 25, WHITE, True)
    sub = doc.add_paragraph()
    set_paragraph(sub, 0, 20, 1.1, WD_ALIGN_PARAGRAPH.CENTER)
    shade(sub._p, BLACK)
    set_run(sub.add_run(subtitle), 13, accent, True)

    add_text(doc, '{{proposta_titulo}}', size=18, color=INK, bold=True,
             before=25, after=5, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
    add_text(doc, 'Preparada especialmente para', size=9.5, color=MUTED,
             before=0, after=3, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
    add_text(doc, '{{cliente_nome}}', size=16, color=accent, bold=True,
             before=0, after=22, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
    add_label_value_table(doc, [
        ('Proposta', '{{proposta_numero}}'),
        ('Data de emissão', '{{proposta_data}}'),
        ('Validade', '{{proposta_validade}}'),
        ('Responsável comercial', '{{responsavel_nome}}'),
    ], accent)
    add_callout(doc, 'OBJETIVO', '{{resumo_executivo_da_solucao}}', accent)


def add_acceptance(doc, accent):
    add_heading(doc, 'Aceite da proposta', 1, accent)
    add_text(doc,
             'Ao aprovar esta proposta, o cliente confirma o interesse na contratação conforme o escopo e as condições comerciais apresentados. O contrato definitivo será gerado pelo aplicativo para revisão e assinatura das partes.',
             size=10.2, color=INK, before=0, after=10)
    add_label_value_table(doc, [
        ('Nome / Razão social', '{{aceite_nome}}'),
        ('CPF / CNPJ', '{{aceite_documento}}'),
        ('Responsável', '{{aceite_responsavel}}'),
        ('Data', '{{aceite_data}}'),
        ('Assinatura', '____________________________________________'),
    ], accent)
    add_callout(doc, 'PRÓXIMO PASSO', 'Após o aceite, o sistema gerará o contrato correspondente e encaminhará o projeto para planejamento, agenda e financeiro.', accent)


def build_plugin():
    doc = Document()
    configure_doc(doc, 'Plug In Eventos', ORANGE)
    add_cover(doc, PLUGIN_PNG, 'Plug In Eventos',
              'Som, iluminação, painéis de LED, DJ e estrutura', ORANGE, 3.15)

    add_page_break(doc)
    add_heading(doc, 'Informações do evento', 1, ORANGE)
    add_label_value_table(doc, [
        ('Cliente', '{{cliente_nome}}'),
        ('Tipo de evento', '{{evento_tipo}}'),
        ('Data e horário', '{{evento_data_horario}}'),
        ('Local', '{{evento_local}}'),
        ('Público estimado', '{{evento_publico}}'),
        ('Montagem / desmontagem', '{{evento_janela_operacional}}'),
    ], ORANGE)

    add_heading(doc, 'Solução proposta', 1, ORANGE)
    add_text(doc, '{{descricao_personalizada_da_solucao}}')
    add_data_table(doc,
        ['Serviço', 'Especificação', 'Quantidade', 'Período'],
        [
            ('Painéis de LED', '{{led_especificacao}}', '{{led_quantidade}}', '{{led_periodo}}'),
            ('Sonorização', '{{som_especificacao}}', '{{som_quantidade}}', '{{som_periodo}}'),
            ('Iluminação', '{{luz_especificacao}}', '{{luz_quantidade}}', '{{luz_periodo}}'),
            ('DJ / operação', '{{dj_especificacao}}', '{{dj_quantidade}}', '{{dj_periodo}}'),
            ('Estruturas', '{{estrutura_especificacao}}', '{{estrutura_quantidade}}', '{{estrutura_periodo}}'),
            ('Serviços adicionais', '{{adicional_especificacao}}', '{{adicional_quantidade}}', '{{adicional_periodo}}'),
        ], [1800, 4800, 1260, 1500], ORANGE)

    # Keep the operational team block together instead of splitting its rows
    # between the solution page and the following page.
    add_page_break(doc)
    add_heading(doc, 'Equipe e operação', 2, ORANGE)
    add_data_table(doc,
        ['Função', 'Quantidade', 'Carga prevista'],
        [
            ('Coordenação técnica', '{{equipe_coordenacao_qtd}}', '{{equipe_coordenacao_horas}}'),
            ('Montagem e desmontagem', '{{equipe_montagem_qtd}}', '{{equipe_montagem_horas}}'),
            ('Operação durante o evento', '{{equipe_operacao_qtd}}', '{{equipe_operacao_horas}}'),
            ('Terceirizados', '{{equipe_terceiros_qtd}}', '{{equipe_terceiros_horas}}'),
        ], [4500, 1800, 3060], ORANGE)
    add_callout(doc, 'PREENCHIMENTO INTELIGENTE', 'A equipe e os equipamentos poderão ser sugeridos pelo sistema com base no tipo de evento e no histórico, permanecendo sujeitos à conferência do responsável.', ORANGE)
    add_heading(doc, 'Premissas operacionais', 2, ORANGE)
    add_bullets(doc, [
        '{{premissa_operacional_1}}',
        '{{premissa_operacional_2}}',
        '{{premissa_operacional_3}}',
        '{{observacao_tecnica_adicional}}',
    ], ORANGE)

    add_page_break(doc)
    add_heading(doc, 'Investimento', 1, ORANGE)
    add_data_table(doc,
        ['Descrição', 'Valor'],
        [
            ('Serviços e equipamentos', '{{valor_servicos_equipamentos}}'),
            ('Equipe e operação', '{{valor_equipe}}'),
            ('Logística e deslocamento', '{{valor_logistica}}'),
            ('Serviços adicionais', '{{valor_adicionais}}'),
            ('Desconto comercial', '{{valor_desconto}}'),
        ], [7000, 2360], ORANGE,
        total_row=('INVESTIMENTO TOTAL', '{{valor_total_proposta}}'))
    add_label_value_table(doc, [
        ('Forma de pagamento', '{{pagamento_forma}}'),
        ('Entrada', '{{pagamento_entrada}}'),
        ('Parcelamento', '{{pagamento_parcelas}}'),
        ('Vencimentos', '{{pagamento_vencimentos}}'),
    ], ORANGE)
    add_callout(doc, 'VALOR PROTEGIDO', 'O aplicativo deverá alertar internamente quando descontos ou alterações reduzirem a margem mínima configurada. Essas informações não aparecerão na proposta enviada ao cliente.', ORANGE)

    add_heading(doc, 'Itens incluídos', 2, ORANGE)
    add_bullets(doc, [
        '{{incluido_1}}', '{{incluido_2}}', '{{incluido_3}}', '{{incluido_4}}'
    ], ORANGE)
    add_heading(doc, 'Itens não incluídos', 2, ORANGE)
    add_bullets(doc, [
        '{{nao_incluido_1}}', '{{nao_incluido_2}}', '{{nao_incluido_3}}'
    ], ORANGE)

    add_heading(doc, 'Responsabilidades do cliente', 2, ORANGE)
    add_bullets(doc, [
        'Garantir acesso ao local nos horários combinados para montagem, testes e desmontagem.',
        'Disponibilizar ponto de energia e condições técnicas conforme orientação prévia.',
        'Informar alterações de horário, local, público ou estrutura com antecedência.',
        '{{responsabilidade_adicional_cliente}}',
    ], ORANGE)

    add_heading(doc, 'Condições comerciais', 2, ORANGE)
    add_bullets(doc, [
        'Esta proposta é válida até {{proposta_validade}}.',
        'A reserva da data ocorre após o aceite e o pagamento definido como entrada.',
        'Mudanças de escopo poderão alterar valores, prazos, equipe e equipamentos.',
        'Cancelamento e reagendamento seguirão as condições do contrato definitivo.',
        '{{condicao_comercial_adicional}}',
    ], ORANGE)
    add_acceptance(doc, ORANGE)

    path = OUT / 'Modelo_Proposta_Plug_In_Eventos.docx'
    doc.save(path)
    return path


def build_led():
    doc = Document()
    configure_doc(doc, 'LED Outdoor', LED_CYAN)
    add_cover(doc, LED_LOGO, 'LED Outdoor',
              'Venda, instalação e configuração de painéis de LED', LED_CYAN, 2.25)

    add_page_break(doc)
    add_heading(doc, 'Dados do projeto', 1, LED_BLUE)
    add_label_value_table(doc, [
        ('Cliente', '{{cliente_nome}}'),
        ('Local de instalação', '{{projeto_local}}'),
        ('Aplicação', '{{projeto_aplicacao}}'),
        ('Ambiente', '{{projeto_ambiente_interno_externo}}'),
        ('Prazo desejado', '{{projeto_prazo_desejado}}'),
        ('Responsável técnico', '{{responsavel_tecnico}}'),
    ], LED_BLUE)

    add_heading(doc, 'Configuração do painel', 1, LED_BLUE)
    add_data_table(doc,
        ['Modelo', 'Largura', 'Altura', 'Área', 'Altura do chão'],
        [('{{painel_modelo}}', '{{painel_largura}}', '{{painel_altura}}', '{{painel_area}}', '{{painel_altura_chao}}')],
        [2200, 1600, 1600, 1800, 2160], LED_BLUE)
    add_label_value_table(doc, [
        ('Pixel pitch', '{{painel_pixel_pitch}}'),
        ('Brilho', '{{painel_brilho}}'),
        ('Resolução estimada', '{{painel_resolucao}}'),
        ('Gabinetes / módulos', '{{painel_quantidade_modulos}}'),
        ('Processadora', '{{painel_processadora}}'),
        ('Software', '{{painel_software}}'),
    ], LED_BLUE)
    add_callout(doc, 'PROJETO VISUAL', '{{descricao_da_disposicao_do_painel_e_estrutura}}', LED_MAGENTA)
    add_text(doc, '[ESPAÇO RESERVADO PARA IMAGEM, RENDER OU FOTOMONTAGEM DO PROJETO]',
             size=10, color=MUTED, bold=True, before=14, after=18,
             align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)

    # Start the supply scope on a clean page so the table is not orphaned.
    add_page_break(doc)
    add_heading(doc, 'Itens e serviços fornecidos', 1, LED_BLUE)
    add_data_table(doc,
        ['Item', 'Descrição', 'Quantidade'],
        [
            ('Painel de LED', '{{fornecimento_painel}}', '{{qtd_painel}}'),
            ('Processamento', '{{fornecimento_processadora}}', '{{qtd_processadora}}'),
            ('Estrutura / serralheria', '{{fornecimento_estrutura}}', '{{qtd_estrutura}}'),
            ('Fundação', '{{fornecimento_fundacao}}', '{{qtd_fundacao}}'),
            ('Instalação elétrica', '{{fornecimento_eletrica}}', '{{qtd_eletrica}}'),
            ('Software e configuração', '{{fornecimento_software}}', '{{qtd_software}}'),
            ('Treinamento', '{{fornecimento_treinamento}}', '{{qtd_treinamento}}'),
            ('Frete', '{{fornecimento_frete}}', '{{qtd_frete}}'),
        ], [2300, 5260, 1800], LED_BLUE)

    add_page_break(doc)
    add_heading(doc, 'Investimento', 1, LED_BLUE)
    add_data_table(doc,
        ['Descrição', 'Valor'],
        [
            ('Painel e componentes', '{{valor_painel_componentes}}'),
            ('Processadora e software', '{{valor_processamento}}'),
            ('Estrutura e fundação', '{{valor_estrutura_fundacao}}'),
            ('Instalação e configuração', '{{valor_instalacao_configuracao}}'),
            ('Frete e deslocamento', '{{valor_frete_deslocamento}}'),
            ('Itens opcionais', '{{valor_opcionais}}'),
            ('Desconto comercial', '{{valor_desconto}}'),
        ], [7000, 2360], LED_BLUE,
        total_row=('INVESTIMENTO TOTAL', '{{valor_total_proposta}}'))
    add_label_value_table(doc, [
        ('Pagamento à vista', '{{pagamento_avista}}'),
        ('Entrada', '{{pagamento_entrada}}'),
        ('Parcelamento', '{{pagamento_parcelas}}'),
        ('Prazo de entrega', '{{prazo_entrega}}'),
    ], LED_BLUE)

    add_heading(doc, 'Cronograma previsto', 2, LED_BLUE)
    add_data_table(doc,
        ['Etapa', 'Prazo'],
        [
            ('Aprovação e contrato', '{{cronograma_aprovacao}}'),
            ('Produção / importação', '{{cronograma_producao}}'),
            ('Preparação do local', '{{cronograma_preparacao}}'),
            ('Instalação e testes', '{{cronograma_instalacao}}'),
            ('Entrega e treinamento', '{{cronograma_entrega}}'),
        ], [6200, 3160], LED_BLUE)

    add_heading(doc, 'Garantia e suporte', 2, LED_BLUE)
    add_bullets(doc, [
        'Garantia dos equipamentos: {{garantia_equipamentos}}.',
        'Garantia da instalação: {{garantia_instalacao}}.',
        'Prazo de atendimento técnico: {{suporte_prazo}}.',
        'Plano de manutenção opcional: {{manutencao_opcional}}.',
        '{{garantia_observacao_adicional}}',
    ], LED_BLUE)

    add_heading(doc, 'Responsabilidades do cliente', 2, LED_BLUE)
    add_bullets(doc, [
        'Disponibilizar acesso ao local para vistoria, instalação e testes.',
        'Providenciar autorizações, licenças ou liberações locais quando aplicáveis.',
        'Garantir infraestrutura elétrica e de dados conforme especificação técnica.',
        'Confirmar medidas e condições do local antes do início da produção.',
        '{{responsabilidade_adicional_cliente}}',
    ], LED_BLUE)

    add_heading(doc, 'Condições comerciais', 2, LED_BLUE)
    add_bullets(doc, [
        'Esta proposta é válida até {{proposta_validade}}.',
        'O prazo de entrega começa após assinatura do contrato e confirmação do pagamento inicial.',
        'Alterações no projeto poderão gerar revisão de preço e prazo.',
        'Serviços civis, elétricos ou de conectividade não descritos expressamente não estão incluídos.',
        '{{condicao_comercial_adicional}}',
    ], LED_BLUE)
    add_acceptance(doc, LED_BLUE)

    path = OUT / 'Modelo_Proposta_LED_Outdoor.docx'
    doc.save(path)
    return path


if __name__ == '__main__':
    paths = [build_plugin(), build_led()]
    for path in paths:
        print(path)
